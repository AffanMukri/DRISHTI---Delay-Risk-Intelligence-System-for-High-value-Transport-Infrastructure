from __future__ import annotations

import csv
import io
import json
import math
import re
import statistics
import zipfile
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import load_workbook

from app.errors import AppError


REQUIRED_FIELDS = ("project_code", "reporting_month")
NUMERIC_FIELDS = {
    "approved_cost", "revised_cost", "expenditure", "physical_progress",
    "planned_progress", "financial_progress", "delay_days", "milestones_total",
    "milestones_completed", "milestones_delayed", "milestones_at_risk",
    "land_acquisition_target", "land_acquisition_completed", "land_acquisition_progress",
}
INTEGER_FIELDS = {
    "delay_days", "milestones_total", "milestones_completed", "milestones_delayed",
    "milestones_at_risk",
}
PERCENTAGE_FIELDS = {
    "physical_progress", "planned_progress", "financial_progress", "land_acquisition_progress",
}
DATE_FIELDS = {
    "reporting_month", "original_completion_date", "revised_completion_date",
    "forecast_completion_date",
}

FIELD_ALIASES: dict[str, set[str]] = {
    "project_code": {"project_code", "project_id", "projectid", "cuf_project_id", "code", "project_no", "project_number"},
    "reporting_month": {"reporting_month", "report_month", "month", "monitoring_month", "cuf_month", "period"},
    "approved_cost": {"approved_cost", "original_cost", "sanctioned_cost", "approved_project_cost"},
    "revised_cost": {"revised_cost", "latest_cost", "anticipated_cost", "current_cost"},
    "expenditure": {"expenditure", "cumulative_expenditure", "expenditure_to_date", "actual_expenditure"},
    "physical_progress": {"physical_progress", "actual_progress", "progress", "physical_progress_percent"},
    "planned_progress": {"planned_progress", "scheduled_progress", "target_progress", "expected_progress"},
    "financial_progress": {"financial_progress", "financial_progress_percent"},
    "original_completion_date": {"original_completion_date", "original_cod", "scheduled_completion", "original_completion"},
    "revised_completion_date": {"revised_completion_date", "revised_cod", "revised_completion"},
    "forecast_completion_date": {"forecast_completion_date", "forecast_cod", "expected_completion_date"},
    "delay_days": {"delay_days", "time_overrun_days", "schedule_delay_days"},
    "milestones_total": {"milestones_total", "total_milestones"},
    "milestones_completed": {"milestones_completed", "completed_milestones"},
    "milestones_delayed": {"milestones_delayed", "delayed_milestones"},
    "milestones_at_risk": {"milestones_at_risk", "at_risk_milestones"},
    "land_acquisition_target": {"land_acquisition_target", "land_target", "land_required"},
    "land_acquisition_completed": {"land_acquisition_completed", "land_acquired", "land_completed"},
    "land_acquisition_unit": {"land_acquisition_unit", "land_unit"},
    "land_acquisition_progress": {"land_acquisition_progress", "land_progress", "land_acquisition_percent"},
    "clearance_status": {"clearance_status", "clearances", "statutory_clearances"},
    "contract_status": {"contract_status", "contract_stage", "contract_state"},
    "issues": {"issues", "key_issues", "constraints", "bottlenecks"},
    "remarks": {"remarks", "comments", "observations", "notes"},
}


def _header_key(value: Any) -> str:
    return re.sub(r"_+", "_", re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower())).strip("_")


ALIAS_LOOKUP = {alias: canonical for canonical, aliases in FIELD_ALIASES.items() for alias in aliases}


@dataclass(slots=True)
class ParsedTable:
    columns: list[str]
    rows: list[dict[str, Any]]


@dataclass(slots=True)
class ValidationResult:
    columns: list[str]
    mapping: dict[str, str | None]
    rows: list[dict[str, Any]]
    summary: dict[str, Any]


def _clean_cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (date, datetime, int, float, bool)):
        return value
    text = str(value).strip()
    return text if text else None


def parse_cuf_file(content: bytes, file_type: str, max_rows: int) -> ParsedTable:
    if file_type == "csv":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise AppError("CSV files must use UTF-8 encoding.", code="invalid_file_encoding") from exc
        if "\x00" in text:
            raise AppError("The CSV contains invalid null bytes.", code="invalid_file")
        sample = text[:8192]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(text), dialect)
        records = list(reader)
    else:
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                uncompressed = sum(item.file_size for item in archive.infolist())
                if uncompressed > 50 * 1024 * 1024:
                    raise AppError("The XLSX expands beyond the safe processing limit.", code="file_too_large", status_code=413)
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheet = workbook.active
            records = [list(row) for row in sheet.iter_rows(values_only=True)]
            workbook.close()
        except AppError:
            raise
        except (OSError, ValueError, zipfile.BadZipFile) as exc:
            raise AppError("The XLSX file is invalid or unreadable.", code="invalid_file") from exc

    while records and not any(_clean_cell(value) is not None for value in records[-1]):
        records.pop()
    if not records:
        raise AppError("The uploaded file is empty.", code="empty_file")

    columns = [str(value or "").strip() for value in records[0]]
    if not columns or any(not column for column in columns):
        raise AppError("Every source column must have a non-empty header.", code="invalid_headers")
    normalized_headers = [_header_key(column) for column in columns]
    duplicates = [name for name, count in Counter(normalized_headers).items() if count > 1]
    if duplicates:
        raise AppError("Duplicate source column names were detected.", code="duplicate_headers", details=duplicates)

    data_rows = records[1:]
    if len(data_rows) > max_rows:
        raise AppError(f"The file exceeds the {max_rows:,}-row limit.", code="too_many_rows", status_code=413)

    rows: list[dict[str, Any]] = []
    for values in data_rows:
        padded = list(values) + [None] * max(0, len(columns) - len(values))
        row = {column: _clean_cell(padded[index]) for index, column in enumerate(columns)}
        if any(value is not None for value in row.values()):
            rows.append(row)
    if not rows:
        raise AppError("The uploaded file contains headers but no data rows.", code="empty_file")
    return ParsedTable(columns=columns, rows=rows)


def detect_mapping(columns: list[str]) -> dict[str, str | None]:
    mapping: dict[str, str | None] = {}
    claimed: set[str] = set()
    for column in columns:
        canonical = ALIAS_LOOKUP.get(_header_key(column))
        if canonical in claimed:
            canonical = None
        mapping[column] = canonical
        if canonical:
            claimed.add(canonical)
    return mapping


def _is_missing(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _parse_date(value: Any, field: str) -> date:
    if isinstance(value, datetime):
        parsed = value.date()
    elif isinstance(value, date):
        parsed = value
    else:
        text = str(value).strip()
        parsed = None
        for pattern in ("%Y-%m-%d", "%Y-%m", "%d-%m-%Y", "%d/%m/%Y", "%Y/%m/%d", "%m/%Y", "%b %Y", "%B %Y"):
            try:
                parsed = datetime.strptime(text, pattern).date()
                break
            except ValueError:
                continue
        if parsed is None:
            raise ValueError("must be a recognized government reporting date")
    return parsed.replace(day=1) if field == "reporting_month" else parsed


def _parse_number(value: Any, integer: bool) -> int | float:
    if isinstance(value, bool):
        raise ValueError("must be numeric")
    if isinstance(value, (int, float, Decimal)):
        decimal = Decimal(str(value))
    else:
        text = str(value).strip()
        negative = text.startswith("(") and text.endswith(")")
        text = re.sub(r"(?i)\b(?:inr|rs\.?|crores?|cr\.?)\b", "", text)
        text = text.replace("₹", "").replace(",", "").replace("%", "").strip()
        if negative:
            text = f"-{text[1:-1]}"
        try:
            decimal = Decimal(text)
        except InvalidOperation as exc:
            raise ValueError("must be numeric") from exc
    if not decimal.is_finite():
        raise ValueError("must be a finite number")
    if integer:
        if decimal != decimal.to_integral_value():
            raise ValueError("must be a whole number")
        return int(decimal)
    return float(decimal)


def _normalize_complex(field: str, value: Any) -> Any:
    if field == "issues":
        if isinstance(value, list):
            return [str(item).strip() for item in value if str(item).strip()]
        return [part.strip() for part in re.split(r"[;|]", str(value)) if part.strip()]
    if field == "clearance_status":
        if isinstance(value, dict):
            return value
        text = str(value).strip()
        if text.startswith("{"):
            parsed = json.loads(text)
            if not isinstance(parsed, dict):
                raise ValueError("must be a JSON object or semicolon-separated key:value list")
            return parsed
        result: dict[str, str] = {}
        for part in text.split(";"):
            if not part.strip():
                continue
            if ":" not in part:
                raise ValueError("must use key:value pairs separated by semicolons")
            key, item = part.split(":", 1)
            result[key.strip()] = item.strip()
        return result
    if field == "contract_status":
        category = _header_key(value)
        aliases = {
            "active": "active",
            "awarded": "active",
            "ongoing": "in_progress",
            "in_progress": "in_progress",
            "under_execution": "in_progress",
            "pre_award": "pre_award",
            "tendering": "pre_award",
            "completed": "completed",
            "closed": "completed",
            "terminated": "terminated",
        }
        return aliases.get(category, category)
    if field == "land_acquisition_unit":
        unit = _header_key(value)
        return {"hectare": "ha", "hectares": "ha", "acre": "acres"}.get(unit, unit)
    return re.sub(r"\s+", " ", str(value).strip())


def _json_value(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    return value


def validate_rows(
    table: ParsedTable,
    project_lookup: dict[str, tuple[str, str]],
    existing_keys: set[tuple[str, date]],
) -> ValidationResult:
    mapping = detect_mapping(table.columns)
    mapped_fields = {canonical for canonical in mapping.values() if canonical}
    missing_required_columns = [field for field in REQUIRED_FIELDS if field not in mapped_fields]
    output_rows: list[dict[str, Any]] = []
    seen_keys: set[tuple[str, date]] = set()
    error_counts: Counter[str] = Counter()
    transformation_count = 0

    for offset, raw in enumerate(table.rows, start=2):
        normalized: dict[str, Any] = {}
        errors: list[dict[str, Any]] = []
        warnings: list[dict[str, Any]] = []
        transformations: list[dict[str, Any]] = []
        missing_count = 0

        if missing_required_columns:
            for field in missing_required_columns:
                errors.append({"field": field, "code": "missing_required_column", "message": f"No source column maps to required field '{field}'."})

        for source, canonical in mapping.items():
            if not canonical:
                continue
            original = raw.get(source)
            if _is_missing(original):
                missing_count += 1
                if canonical in REQUIRED_FIELDS:
                    errors.append({"field": canonical, "code": "required_value_missing", "message": "A required value is missing."})
                continue
            try:
                if canonical in DATE_FIELDS:
                    value = _parse_date(original, canonical)
                elif canonical in NUMERIC_FIELDS:
                    value = _parse_number(original, canonical in INTEGER_FIELDS)
                elif canonical in {"issues", "clearance_status"}:
                    value = _normalize_complex(canonical, original)
                else:
                    value = _normalize_complex(canonical, original)
                normalized[canonical] = _json_value(value)
                if str(_json_value(original)) != str(normalized[canonical]):
                    transformations.append({
                        "field": canonical,
                        "original": _json_value(original),
                        "normalized": normalized[canonical],
                        "rule": "canonical_format",
                    })
            except (ValueError, TypeError, json.JSONDecodeError) as exc:
                errors.append({"field": canonical, "code": "invalid_value", "message": str(exc), "value": _json_value(original)})

        for field in ("approved_cost", "revised_cost", "expenditure", "land_acquisition_target", "land_acquisition_completed"):
            if field in normalized and normalized[field] < 0:
                errors.append({"field": field, "code": "out_of_range", "message": "Value cannot be negative."})
        for field in PERCENTAGE_FIELDS:
            if field in normalized and not 0 <= normalized[field] <= 100:
                errors.append({"field": field, "code": "out_of_range", "message": "Percentage must be between 0 and 100."})
        for field in INTEGER_FIELDS - {"delay_days"}:
            if field in normalized and normalized[field] < 0:
                errors.append({"field": field, "code": "out_of_range", "message": "Count cannot be negative."})

        total = normalized.get("milestones_total")
        parts = sum(normalized.get(field, 0) for field in ("milestones_completed", "milestones_delayed", "milestones_at_risk"))
        if total is not None and parts > total:
            errors.append({"field": "milestones_total", "code": "inconsistent_value", "message": "Milestone status counts exceed total milestones."})

        project_database_id = None
        project_code = normalized.get("project_code")
        reporting_value = normalized.get("reporting_month")
        reporting_month = date.fromisoformat(reporting_value) if reporting_value else None
        if project_code:
            project_match = project_lookup.get(str(project_code).casefold())
            if project_match:
                project_database_id, certified_code = project_match
                if certified_code != project_code:
                    transformations.append({"field": "project_code", "original": project_code, "normalized": certified_code, "rule": "certified_project_code"})
                    normalized["project_code"] = certified_code
                    project_code = certified_code
            else:
                errors.append({"field": "project_code", "code": "unknown_project", "message": f"Project '{project_code}' does not exist."})

        is_duplicate = False
        if project_database_id and reporting_month:
            key = (project_database_id, reporting_month)
            if key in existing_keys:
                is_duplicate = True
                errors.append({"field": "reporting_month", "code": "duplicate_monthly_import", "message": "A certified CUF update already exists for this project and month."})
            elif key in seen_keys:
                is_duplicate = True
                errors.append({"field": "reporting_month", "code": "duplicate_file_row", "message": "This project and reporting month appears more than once in the file."})
            else:
                seen_keys.add(key)

        for error in errors:
            error_counts[error["code"]] += 1
        transformation_count += len(transformations)
        output_rows.append({
            "row_number": offset,
            "project_id": project_database_id,
            "project_code": project_code,
            "reporting_month": reporting_month,
            "validation_status": "invalid" if errors else "valid",
            "raw_data": {key: _json_value(value) for key, value in raw.items()},
            "normalized_data": normalized,
            "transformations": transformations,
            "validation_errors": errors,
            "validation_warnings": warnings,
            "missing_value_count": missing_count,
            "is_duplicate": is_duplicate,
            "anomaly_count": 0,
        })

    # IQR-based warnings are informational. They never rewrite or reject source data.
    for field in sorted(NUMERIC_FIELDS - INTEGER_FIELDS):
        candidates = [float(row["normalized_data"][field]) for row in output_rows if row["validation_status"] == "valid" and field in row["normalized_data"]]
        if len(candidates) < 4:
            continue
        quartiles = statistics.quantiles(candidates, n=4, method="inclusive")
        lower, upper = quartiles[0], quartiles[2]
        iqr = upper - lower
        if math.isclose(iqr, 0):
            continue
        floor, ceiling = lower - 1.5 * iqr, upper + 1.5 * iqr
        for row in output_rows:
            value = row["normalized_data"].get(field)
            if row["validation_status"] == "valid" and value is not None and (value < floor or value > ceiling):
                row["validation_warnings"].append({
                    "field": field,
                    "code": "statistical_outlier",
                    "message": f"Value is outside the IQR range ({floor:.2f} to {ceiling:.2f}).",
                    "value": value,
                })
                row["anomaly_count"] += 1

    total_rows = len(output_rows)
    valid_rows = sum(row["validation_status"] == "valid" for row in output_rows)
    invalid_rows = total_rows - valid_rows
    missing_values = sum(row["missing_value_count"] for row in output_rows)
    duplicate_rows = sum(row["is_duplicate"] for row in output_rows)
    anomaly_rows = sum(row["anomaly_count"] > 0 for row in output_rows)
    mapped_cell_count = max(1, total_rows * max(1, len(mapped_fields)))
    quality_score = round(max(0.0, min(100.0,
        60 * (valid_rows / total_rows)
        + 15 * (1 - min(1, missing_values / mapped_cell_count))
        + 15 * (1 - duplicate_rows / total_rows)
        + 10 * (1 - anomaly_rows / total_rows)
    )), 2)
    summary = {
        "total_rows": total_rows,
        "valid_rows": valid_rows,
        "invalid_rows": invalid_rows,
        "missing_values": missing_values,
        "duplicate_rows": duplicate_rows,
        "anomaly_rows": anomaly_rows,
        "quality_score": quality_score,
        "required_fields": list(REQUIRED_FIELDS),
        "missing_required_columns": missing_required_columns,
        "unmapped_columns": [column for column, canonical in mapping.items() if not canonical],
        "transformation_count": transformation_count,
        "error_counts": dict(error_counts),
        "quality_formula": "60% valid rows + 15% completeness + 15% uniqueness + 10% anomaly-free rows",
    }
    return ValidationResult(columns=table.columns, mapping=mapping, rows=output_rows, summary=summary)
