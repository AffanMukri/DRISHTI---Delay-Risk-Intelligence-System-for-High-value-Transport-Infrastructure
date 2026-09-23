from __future__ import annotations

import csv
import io
import json
from datetime import date, datetime
from typing import Any, Iterable
from uuid import UUID
from xml.sax.saxutils import escape

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    KeepTogether,
    LongTable,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


NAVY = colors.HexColor("#17365D")
BLUE = colors.HexColor("#DCE6F1")
LIGHT = colors.HexColor("#F3F6FA")
RED = colors.HexColor("#C00000")
AMBER = colors.HexColor("#F4B183")
GREEN = colors.HexColor("#70AD47")
GRID = colors.HexColor("#CBD5E1")


def _display(value: Any, *, number: bool = False) -> str:
    if value is None:
        return "Not available"
    if isinstance(value, (date, datetime)):
        return value.strftime("%d %b %Y")
    if isinstance(value, float):
        return f"{value:,.1f}" if number else f"{value:.1f}"
    return str(value)


def _paragraph(value: Any, style: ParagraphStyle) -> Paragraph:
    return Paragraph(escape(_display(value)), style)


def _footer(canvas: Any, document: Any) -> None:
    canvas.saveState()
    canvas.setStrokeColor(GRID)
    canvas.line(18 * mm, 13 * mm, document.pagesize[0] - 18 * mm, 13 * mm)
    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(colors.HexColor("#64748B"))
    canvas.drawString(18 * mm, 8 * mm, "DHRISTI - generated from authenticated project monitoring records")
    canvas.drawRightString(document.pagesize[0] - 18 * mm, 8 * mm, f"Page {canvas.getPageNumber()}")
    canvas.restoreState()


def build_pdf(report: dict[str, Any]) -> bytes:
    buffer = io.BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=16 * mm,
        leftMargin=16 * mm,
        topMargin=14 * mm,
        bottomMargin=18 * mm,
        title=report["title"],
        author="DHRISTI",
        subject=f"{report['title']} for {report['reporting_month']:%B %Y}",
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="PXTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=18, leading=22, textColor=NAVY, alignment=TA_CENTER, spaceAfter=4))
    styles.add(ParagraphStyle(name="PXSub", parent=styles["Normal"], fontName="Helvetica", fontSize=8.5, leading=11, textColor=colors.HexColor("#475569"), alignment=TA_CENTER))
    styles.add(ParagraphStyle(name="PXHeading", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=11, leading=14, textColor=NAVY, spaceBefore=8, spaceAfter=5))
    styles.add(ParagraphStyle(name="PXBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=8.5, leading=11, textColor=colors.HexColor("#1E293B"), alignment=TA_LEFT))
    styles.add(ParagraphStyle(name="PXSmall", parent=styles["BodyText"], fontName="Helvetica", fontSize=7, leading=9, textColor=colors.HexColor("#334155")))
    styles.add(ParagraphStyle(name="PXCell", parent=styles["BodyText"], fontName="Helvetica", fontSize=6.8, leading=8.2, textColor=colors.HexColor("#1E293B")))
    styles.add(ParagraphStyle(name="PXHeader", parent=styles["BodyText"], fontName="Helvetica-Bold", fontSize=6.7, leading=8, textColor=colors.white, alignment=TA_CENTER))

    story: list[Any] = [
        Paragraph("GOVERNMENT OF INDIA", styles["PXSub"]),
        Paragraph("MINISTRY OF STATISTICS AND PROGRAMME IMPLEMENTATION", styles["PXTitle"]),
        Paragraph(report["title"], styles["PXHeading"]),
        Paragraph(
            f"Reference period: {report['reporting_month']:%B %Y} | Data cut-off: {report['data_as_of']:%d %b %Y} | Report ID: {escape(str(report['report_id']))}",
            styles["PXSub"],
        ),
        Spacer(1, 5 * mm),
    ]

    summary = report["summary"]
    kpis = [
        ("Projects", summary["total_projects"]),
        ("Original cost (INR Cr)", summary["original_approved_cost"]),
        ("Revised cost (INR Cr)", summary["latest_revised_cost"]),
        ("Cost escalation", f"{summary['cost_escalation_percentage']:.1f}%" if summary["cost_escalation_percentage"] is not None else "Not available"),
        ("Delayed projects", summary["delayed_projects"]),
        ("High / Critical", f"{summary['high_risk_projects']} / {summary['critical_projects']}"),
        ("Open warnings", summary["open_warnings"]),
        ("Open interventions", summary["open_interventions"] if summary["intervention_data_available"] else "Restricted / unavailable"),
    ]
    kpi_table = Table(
        [[_paragraph(label, styles["PXSmall"]) for label, _ in kpis], [_paragraph(_display(value, number=True), styles["PXBody"]) for _, value in kpis]],
        colWidths=[(landscape(A4)[0] - 32 * mm) / len(kpis)] * len(kpis),
    )
    kpi_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), BLUE),
        ("BACKGROUND", (0, 1), (-1, 1), LIGHT),
        ("TEXTCOLOR", (0, 1), (-1, 1), NAVY),
        ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("BOX", (0, 0), (-1, -1), 0.5, GRID),
        ("INNERGRID", (0, 0), (-1, -1), 0.25, GRID),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.extend([kpi_table, Paragraph("Executive summary", styles["PXHeading"]), Paragraph(escape(report["executive_summary"]), styles["PXBody"])])

    project_rows = report["projects"]
    if project_rows:
        story.append(Paragraph("Project performance", styles["PXHeading"]))
        headers = ["Project", "Ministry / Sector", "Original cost", "Revised cost", "Escalation", "Delay", "Progress", "Risk"]
        data = [[_paragraph(header, styles["PXHeader"]) for header in headers]]
        for row in project_rows:
            risk_label = str(row.get("risk_level") or "Not assessed").replace("_", " ").title()
            data.append([
                _paragraph(f"{row['project_id']} - {row['project_name']}", styles["PXCell"]),
                _paragraph(f"{row['ministry']} / {row['sector']}", styles["PXCell"]),
                _paragraph(_display(row.get("original_approved_cost"), number=True), styles["PXCell"]),
                _paragraph(_display(row.get("latest_revised_cost"), number=True), styles["PXCell"]),
                _paragraph(f"{row['cost_escalation_percentage']:.1f}%" if row.get("cost_escalation_percentage") is not None else "Not available", styles["PXCell"]),
                _paragraph(f"{row['schedule_delay_days']} days" if row.get("schedule_delay_days") is not None else "Not available", styles["PXCell"]),
                _paragraph(f"{row['physical_progress']:.1f}%" if row.get("physical_progress") is not None else "Not available", styles["PXCell"]),
                _paragraph(f"{risk_label} ({_display(row.get('overall_risk_score'))})", styles["PXCell"]),
            ])
        table = LongTable(data, repeatRows=1, colWidths=[54 * mm, 45 * mm, 25 * mm, 25 * mm, 22 * mm, 21 * mm, 20 * mm, 27 * mm])
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), NAVY),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("GRID", (0, 0), (-1, -1), 0.25, GRID),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LIGHT]),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(table)
    else:
        story.extend([Paragraph("Project performance", styles["PXHeading"]), Paragraph("No project records matched this report scope and period.", styles["PXBody"])])

    if report["warnings"]:
        story.append(Paragraph("Emerging and unresolved warnings", styles["PXHeading"]))
        data = [[_paragraph(value, styles["PXHeader"]) for value in ["Warning", "Project", "Severity", "Trigger / evidence", "Detected", "Recommended action"]]]
        for row in report["warnings"]:
            evidence = row.get("evidence") or []
            evidence_text = "; ".join(str(item.get("description") or item.get("label") or item) for item in evidence[:2]) if isinstance(evidence, list) else str(evidence)
            data.append([
                _paragraph(f"{row['warning_code']} - {row['title']}", styles["PXCell"]),
                _paragraph(f"{row['project_id']} - {row['project_name']}", styles["PXCell"]),
                _paragraph(str(row["severity"]).title(), styles["PXCell"]),
                _paragraph(evidence_text or row.get("trigger_rule") or "Not recorded", styles["PXCell"]),
                _paragraph(row.get("detected_at"), styles["PXCell"]),
                _paragraph(row.get("recommended_action") or "Not recorded", styles["PXCell"]),
            ])
        warning_table = LongTable(data, repeatRows=1, colWidths=[52 * mm, 47 * mm, 18 * mm, 62 * mm, 25 * mm, 58 * mm])
        warning_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), NAVY), ("GRID", (0, 0), (-1, -1), 0.25, GRID), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LIGHT]), ("PADDING", (0, 0), (-1, -1), 4)]))
        story.append(warning_table)

    if report["interventions"]:
        story.append(Paragraph("Intervention status", styles["PXHeading"]))
        data = [[_paragraph(value, styles["PXHeader"]) for value in ["Intervention", "Project", "Priority", "Status", "Responsible officer", "Due date", "Action"]]]
        for row in report["interventions"]:
            data.append([
                _paragraph(f"{row['intervention_code']} - {row['issue']}", styles["PXCell"]),
                _paragraph(row["project_id"], styles["PXCell"]),
                _paragraph(str(row["priority"]).title(), styles["PXCell"]),
                _paragraph(str(row["status"]).replace("_", " ").title(), styles["PXCell"]),
                _paragraph(row.get("assigned_to_name") or "Unassigned", styles["PXCell"]),
                _paragraph(row.get("due_date"), styles["PXCell"]),
                _paragraph(row["recommended_action"], styles["PXCell"]),
            ])
        intervention_table = LongTable(data, repeatRows=1, colWidths=[51 * mm, 28 * mm, 19 * mm, 22 * mm, 38 * mm, 24 * mm, 80 * mm])
        intervention_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), NAVY), ("GRID", (0, 0), (-1, -1), 0.25, GRID), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LIGHT]), ("PADDING", (0, 0), (-1, -1), 4)]))
        story.append(intervention_table)

    story.append(Paragraph("Analytical recommendations", styles["PXHeading"]))
    recommendation_blocks: list[Any] = []
    for recommendation in report["recommendations"]:
        recommendation_blocks.extend([
            Paragraph(f"<b>{escape(recommendation['title'])}</b> [{escape(recommendation['source_label'])}]", styles["PXBody"]),
            Paragraph(escape(recommendation["recommendation"]), styles["PXSmall"]),
            Paragraph(f"Evidence: {escape(recommendation['evidence'])}", styles["PXSmall"]),
            Spacer(1, 2 * mm),
        ])
    story.append(KeepTogether(recommendation_blocks or [Paragraph("No analytical recommendation was triggered for this scope.", styles["PXBody"])]))
    story.extend([Spacer(1, 4 * mm), Paragraph("Recommendations are deterministic analytical outputs based on the values shown above. They are not administrative approvals or guaranteed project outcomes.", styles["PXSmall"])])

    document.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return buffer.getvalue()


def _excel_safe(value: Any) -> Any:
    if isinstance(value, UUID):
        return str(value)
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.astimezone().replace(tzinfo=None)
    if isinstance(value, (list, dict)):
        value = json.dumps(value, ensure_ascii=False, default=str)
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _write_sheet(workbook: Workbook, name: str, rows: Iterable[dict[str, Any]]) -> None:
    sheet = workbook.create_sheet(name)
    values = list(rows)
    if not values:
        sheet["A1"] = "No records available for this report scope."
        sheet["A1"].font = Font(italic=True, color="64748B")
        sheet.sheet_view.showGridLines = False
        return
    headers = list(values[0].keys())
    sheet.append([header.replace("_", " ").title() for header in headers])
    for row in values:
        sheet.append([_excel_safe(row.get(header)) for header in headers])
    header_fill = PatternFill("solid", fgColor="17365D")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=9, color="1E293B")
            cell.alignment = Alignment(vertical="top", wrap_text=False)
            if isinstance(cell.value, (date, datetime)):
                cell.number_format = "dd-mmm-yyyy"
            elif isinstance(cell.value, float):
                cell.number_format = "#,##0.0"
            elif isinstance(cell.value, int):
                cell.number_format = "#,##0"
    for index, header in enumerate(headers, start=1):
        max_length = max(len(str(header)), *(len(str(row.get(header) or "")) for row in values))
        sheet.column_dimensions[get_column_letter(index)].width = min(max(max_length + 2, 12), 48)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False


def build_xlsx(report: dict[str, Any]) -> bytes:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    summary_sheet.sheet_view.showGridLines = False
    summary_sheet["A2"] = report["title"]
    summary_sheet["A2"].font = Font(name="Arial", size=15, bold=True, color="17365D")
    summary_sheet["A3"] = f"Reference period: {report['reporting_month']:%B %Y}"
    summary_sheet["A4"] = f"Data cut-off: {report['data_as_of']:%d %b %Y}"
    summary_sheet["A5"] = f"Report ID: {report['report_id']}"
    summary_sheet["A7"] = "Portfolio metric"
    summary_sheet["B7"] = "Value"
    summary_sheet["C7"] = "Definition"
    definitions = {
        "total_projects": "Projects matching the selected scope.",
        "original_approved_cost": "Sum of earliest reported approved costs (INR crore).",
        "latest_revised_cost": "Sum of latest revised costs at the reporting cut-off (INR crore).",
        "cumulative_expenditure": "Sum of latest cumulative expenditure at the reporting cut-off (INR crore).",
        "absolute_cost_escalation": "Latest revised cost less original approved cost (INR crore).",
        "cost_escalation_percentage": "Portfolio escalation divided by original approved cost.",
        "delayed_projects": "Projects with a positive schedule delay.",
        "high_risk_projects": "Projects whose latest available risk level is High.",
        "critical_projects": "Projects whose latest available risk level is Critical.",
        "open_warnings": "Warnings unresolved at the reporting cut-off.",
        "open_interventions": "Interventions unresolved at the reporting cut-off.",
    }
    row_index = 8
    for key, definition in definitions.items():
        summary_sheet.cell(row=row_index, column=1, value=key.replace("_", " ").title())
        summary_sheet.cell(row=row_index, column=2, value=_excel_safe(report["summary"].get(key)))
        summary_sheet.cell(row=row_index, column=3, value=definition)
        row_index += 1
    row_index += 1
    summary_sheet.cell(row=row_index, column=1, value="Executive summary")
    summary_sheet.cell(row=row_index + 1, column=1, value=report["executive_summary"])
    summary_sheet.merge_cells(start_row=row_index + 1, start_column=1, end_row=row_index + 2, end_column=3)
    summary_sheet.cell(row=row_index + 1, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    summary_sheet.row_dimensions[row_index + 1].height = 34
    for cell in summary_sheet[7]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    for row in summary_sheet.iter_rows(min_row=2, max_row=summary_sheet.max_row, min_col=1, max_col=3):
        for cell in row:
            if cell.font.name != "Arial":
                cell.font = Font(name="Arial", size=10, bold=cell.font.bold, color="1E293B")
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column == 3)
    summary_sheet.column_dimensions["A"].width = 34
    summary_sheet.column_dimensions["B"].width = 22
    summary_sheet.column_dimensions["C"].width = 74
    summary_sheet.freeze_panes = "A8"
    summary_sheet.sheet_properties.tabColor = "17365D"

    _write_sheet(workbook, "Projects", report["projects"])
    _write_sheet(workbook, "Warnings", report["warnings"])
    _write_sheet(workbook, "Interventions", report["interventions"])
    _write_sheet(workbook, "Risk Drivers", report["risk_drivers"])
    _write_sheet(workbook, "Recommendations", report["recommendations"])
    output = io.BytesIO()
    workbook.save(output)
    # Re-open once so malformed workbook relationships fail before delivery.
    load_workbook(io.BytesIO(output.getvalue()), read_only=True, data_only=False).close()
    return output.getvalue()


def build_csv(report: dict[str, Any]) -> bytes:
    rows = report["interventions"] if report["report_type"] == "intervention" else report["projects"]
    output = io.StringIO(newline="")
    if not rows:
        output.write("message\r\nNo records available for this report scope.\r\n")
    else:
        headers = list(rows[0].keys())
        writer = csv.DictWriter(output, fieldnames=headers, extrasaction="ignore", lineterminator="\r\n")
        writer.writeheader()
        for row in rows:
            writer.writerow({key: _excel_safe(value) for key, value in row.items()})
    return ("\ufeff" + output.getvalue()).encode("utf-8")
