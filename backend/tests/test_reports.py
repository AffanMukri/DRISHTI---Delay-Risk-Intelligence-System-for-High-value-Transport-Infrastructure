from __future__ import annotations

import csv
import io
from datetime import UTC, date, datetime
from uuid import UUID

from fastapi.testclient import TestClient
import fitz
import pytest
from openpyxl import load_workbook

from app.api.dependencies import get_report_service
from app.auth.dependencies import get_current_profile
from app.auth.models import CurrentProfile
from app.main import create_app
from app.reporting.builders import build_csv, build_pdf, build_xlsx
from app.schemas.report import ReportFormat, ReportGenerateRequest, ReportType
from app.services.reports import GeneratedReport, ReportService


USER_ID = UUID("10000000-0000-0000-0000-000000000001")
REPORT_ID = UUID("90000000-0000-0000-0000-000000000001")
NOW = datetime(2026, 9, 21, 10, 0, tzinfo=UTC)


def assembled_report() -> dict[str, object]:
    project = {
        "project_id": "PX-001",
        "project_name": "National Connectivity Corridor",
        "ministry": "Ministry of Infrastructure",
        "sector": "Transport",
        "original_approved_cost": 1000.0,
        "latest_revised_cost": 1120.0,
        "cumulative_expenditure": 610.0,
        "absolute_cost_escalation": 120.0,
        "cost_escalation_percentage": 12.0,
        "schedule_delay_days": 90,
        "physical_progress": 61.0,
        "overall_risk_score": 72.0,
        "risk_level": "high_risk",
    }
    return {
        "report_id": REPORT_ID,
        "report_type": "monthly_flash",
        "title": "Monthly Flash Report",
        "reporting_month": date(2026, 9, 1),
        "data_as_of": date(2026, 9, 30),
        "summary": {
            "total_projects": 1,
            "original_approved_cost": 1000.0,
            "latest_revised_cost": 1120.0,
            "cumulative_expenditure": 610.0,
            "absolute_cost_escalation": 120.0,
            "cost_escalation_percentage": 12.0,
            "delayed_projects": 1,
            "high_risk_projects": 1,
            "critical_projects": 0,
            "open_warnings": 1,
            "open_interventions": 1,
            "intervention_data_available": True,
        },
        "executive_summary": "One monitored project has a recorded cost and schedule overrun.",
        "projects": [project],
        "warnings": [{
            "warning_code": "WARN-001",
            "title": "Schedule variance threshold exceeded",
            "project_id": "PX-001",
            "project_name": "National Connectivity Corridor",
            "severity": "high",
            "trigger_rule": "Progress variance exceeds threshold",
            "evidence": [{"description": "Actual progress is 9 percentage points behind plan."}],
            "detected_at": NOW,
            "recommended_action": "Review the recovery plan.",
        }],
        "interventions": [{
            "intervention_code": "INT-001",
            "issue": "Schedule variance",
            "project_id": "PX-001",
            "priority": "high",
            "status": "in_progress",
            "assigned_to_name": "Monitoring Officer",
            "due_date": date(2026, 10, 15),
            "recommended_action": "Review the recovery plan.",
        }],
        "risk_drivers": [{"project_id": "PX-001", "name": "Schedule gap", "value": 78.0}],
        "comparisons": [],
        "recommendations": [{
            "title": "Review delivery plan",
            "recommendation": "Confirm accountable recovery actions.",
            "source_type": "analytical_rule",
            "source_label": "Analytical rule output",
            "evidence": "The project is behind schedule.",
        }],
    }


def test_pdf_builder_returns_readable_grounded_document() -> None:
    content = build_pdf(assembled_report())
    assert content.startswith(b"%PDF")
    document = fitz.open(stream=content, filetype="pdf")
    assert document.page_count >= 1
    text = "\n".join(page.get_text() for page in document)
    document.close()
    assert "Monthly Flash Report" in text
    assert "National Connectivity Corridor" in text
    assert "Analytical rule output" in text


def test_xlsx_builder_creates_traceable_sheets_and_typed_values() -> None:
    content = build_xlsx(assembled_report())
    workbook = load_workbook(io.BytesIO(content), data_only=False)
    assert workbook.sheetnames == ["Summary", "Projects", "Warnings", "Interventions", "Risk Drivers", "Recommendations"]
    assert workbook["Projects"]["A2"].value == "PX-001"
    assert isinstance(workbook["Projects"]["E2"].value, (int, float))
    assert workbook["Summary"]["A2"].value == "Monthly Flash Report"
    workbook.close()


def test_csv_builder_uses_real_detail_rows_and_utf8_bom() -> None:
    content = build_csv(assembled_report())
    assert content.startswith(b"\xef\xbb\xbf")
    rows = list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))
    assert rows[0]["project_id"] == "PX-001"
    assert rows[0]["cost_escalation_percentage"] == "12.0"


class FakeRepository:
    def __init__(self) -> None:
        self.export: dict[str, object] | None = None

    async def options(self) -> dict[str, object]:
        return {
            "reporting_months": [date(2026, 9, 1)],
            "sectors": ["Transport"],
            "ministries": [{"value": str(UUID(int=2)), "label": "Ministry of Infrastructure"}],
            "projects": [{"value": "PX-001", "label": "PX-001 - National Connectivity Corridor", "sector": "Transport", "ministry_id": UUID(int=2)}],
        }

    async def project_snapshots(self, **_: object) -> list[dict[str, object]]:
        return [{
            "project_database_id": UUID(int=3),
            "project_id": "PX-001",
            "project_name": "National Connectivity Corridor",
            "ministry_id": UUID(int=2),
            "ministry": "Ministry of Infrastructure",
            "implementing_agency": "National Projects Agency",
            "sector": "Transport",
            "project_type": "Infrastructure",
            "state": "Delhi",
            "project_status": "active",
            "currency": "INR",
            "original_approved_cost": 1000.0,
            "latest_revised_cost": 1120.0,
            "cumulative_expenditure": 610.0,
            "physical_progress": 61.0,
            "planned_progress": 70.0,
            "financial_progress": 54.5,
            "original_completion_date": date(2026, 12, 31),
            "current_completion_date": date(2027, 3, 31),
            "schedule_delay_days": 90,
            "reporting_month": date(2026, 9, 1),
            "overall_risk_score": 72.0,
            "risk_level": "high_risk",
            "cost_risk": 65.0,
            "schedule_risk": 78.0,
            "implementation_risk": 70.0,
            "risk_methodology": "hybrid-v1",
            "risk_assessed_at": NOW,
            "model_version_id": None,
            "risk_drivers": [{"code": "SCHEDULE_GAP", "name": "Schedule gap", "value": 78.0}],
        }]

    async def warnings(self, **_: object) -> list[dict[str, object]]:
        return []

    async def interventions(self, **_: object) -> list[dict[str, object]]:
        return []

    async def record_export(self, **values: object) -> None:
        self.export = values


def profile(role: str = "administrator") -> CurrentProfile:
    return CurrentProfile(id=USER_ID, email="admin@example.com", full_name="Admin", role=role, is_active=True)


@pytest.mark.anyio
async def test_report_service_generates_and_audits_xlsx() -> None:
    repository = FakeRepository()
    service = ReportService(repository)  # type: ignore[arg-type]
    result = await service.generate(ReportGenerateRequest(
        report_type=ReportType.MONTHLY_FLASH,
        output_format=ReportFormat.XLSX,
        reporting_month=date(2026, 9, 1),
    ), profile())
    assert result.content.startswith(b"PK")
    assert result.file_name.endswith(".xlsx")
    assert repository.export is not None
    assert repository.export["checksum_sha256"] == result.checksum_sha256
    assert repository.export["project_count"] == 1


@pytest.mark.anyio
async def test_report_service_hides_intervention_reports_from_analyst() -> None:
    service = ReportService(FakeRepository())  # type: ignore[arg-type]
    options = await service.options(profile("analyst"))
    assert "intervention" not in {item["value"] for item in options["report_types"]}
    assert "pragati_review_dossier" not in {item["value"] for item in options["report_types"]}


class FakeReportService:
    async def options(self, _: CurrentProfile) -> dict[str, object]:
        return {
            "report_types": [{
                "value": "monthly_flash",
                "label": "Monthly Flash Report",
                "description": "Portfolio report",
                "formats": ["pdf", "xlsx"],
                "requires": None,
            }],
            "reporting_months": [date(2026, 9, 1)],
            "sectors": ["Transport"],
            "ministries": [],
            "projects": [],
        }

    async def generate(self, _: ReportGenerateRequest, __: CurrentProfile) -> GeneratedReport:
        return GeneratedReport(
            report_id=REPORT_ID,
            content=b"%PDF-1.4\n%%EOF",
            file_name="dhristi-monthly-flash.pdf",
            media_type="application/pdf",
            data_as_of=date(2026, 9, 30),
            checksum_sha256="a" * 64,
        )


def test_report_endpoint_returns_a_real_attachment_with_trace_headers() -> None:
    app = create_app()
    app.dependency_overrides[get_report_service] = FakeReportService
    app.dependency_overrides[get_current_profile] = lambda: profile()
    with TestClient(app) as client:
        response = client.post("/api/reports/generate", json={
            "reportType": "monthly_flash",
            "outputFormat": "pdf",
            "reportingMonth": "2026-09-01",
        })
    assert response.status_code == 200
    assert response.content.startswith(b"%PDF")
    assert response.headers["content-disposition"].endswith('"dhristi-monthly-flash.pdf"')
    assert response.headers["x-report-id"] == str(REPORT_ID)
    assert response.headers["x-data-as-of"] == "2026-09-30"
